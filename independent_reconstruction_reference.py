import os,glob,sqlite3,datetime,textwrap
import numpy as np,pandas as pd,matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from sklearn.model_selection import train_test_split,StratifiedKFold,cross_val_score
from sklearn.preprocessing import StandardScaler
from sklearn.pipeline import make_pipeline
from sklearn.linear_model import LogisticRegression
from sklearn.svm import SVC
from sklearn.tree import DecisionTreeClassifier
from sklearn.neighbors import KNeighborsClassifier
from sklearn.metrics import accuracy_score,confusion_matrix
P='/mnt/data/spacex_work/IBM-Data-Science-Capstone-SpaceX-main';out='/mnt/data'; chartdir='/mnt/data/spacex_work/charts';os.makedirs(chartdir,exist_ok=True)
a=pd.read_csv(f'{P}/dataset_part_2.csv'); a['Year']=pd.to_datetime(a.Date).dt.year
geo=pd.read_csv(f'{P}/spacex_launch_geo.csv'); dash=pd.read_csv(f'{P}/spacex_launch_dash.csv')
con=sqlite3.connect(f'{P}/my_data1.db')
plt.rcParams.update({'figure.facecolor':'white','axes.facecolor':'white','font.size':12,'axes.titlesize':17,'axes.labelsize':12,'figure.dpi':130})
def save(n): plt.tight_layout(pad=1.8); f=f'{chartdir}/{n}.png';plt.savefig(f,bbox_inches='tight',dpi=150);plt.close();return f
plots={}
def scatter(n,x,y,cat,title,xlabel,ylabel,rotate=0):
 plt.figure(figsize=(11,4.8));
 for k,g in a.groupby(cat): plt.scatter(g[x],g[y],label=str(k),alpha=.74,s=55)
 plt.title(title);plt.xlabel(xlabel);plt.ylabel(ylabel);plt.legend(fontsize=9,ncol=3,loc='best');plt.xticks(rotation=rotate);plt.grid(alpha=.15);plots[n]=save(n)
scatter(18,'LaunchSite','FlightNumber','Class','Flight sequence by site and outcome','Launch site','Flight number',10)
scatter(19,'LaunchSite','PayloadMass','Class','Payload mass, site, and landing outcome','Launch site','Payload mass (kg)',10)
plt.figure(figsize=(11,4.8)); v=a.groupby('Orbit').Class.agg(['mean','count']).sort_values('mean',ascending=False);plt.bar(v.index,v['mean']*100);plt.ylim(0,105);plt.ylabel('Landing success (%)');plt.title('Landing success rate by orbit | sample size labelled');plt.xticks(rotation=25);[plt.text(i,r['mean']*100+2,'n='+str(int(r['count'])),ha='center',fontsize=9) for i,(_,r) in enumerate(v.iterrows())];plots[20]=save(20)
scatter(21,'Orbit','FlightNumber','Class','Launch sequence and orbital destination','Orbit','Flight number',30)
scatter(22,'Orbit','PayloadMass','Class','Payload mass across orbital destinations','Orbit','Payload mass (kg)',30)
plt.figure(figsize=(11,4.8));v=a.groupby('Year').Class.agg(['mean','count']);plt.plot(v.index,v['mean']*100,marker='o',linewidth=2);plt.xticks(v.index);plt.ylim(-5,105);plt.grid(alpha=.25);plt.xlabel('Launch year');plt.ylabel('Landing success (%)');plt.title('Annual booster landing success in 90-launch dataset');plots[23]=save(23)
# geographic plot based on source lat/long, launch outcome markers
plt.figure(figsize=(11,4.8));
for k,g in geo.groupby('class'):
 plt.scatter(g['Long'],g['Lat'],label='Success' if k==1 else 'Non-success',s=45,alpha=.7)
plt.xlabel('Longitude');plt.ylabel('Latitude');plt.title('Falcon 9 launches by launch-site coordinates');plt.grid(alpha=.25);plt.legend();plots[35]=save(35)
plt.figure(figsize=(11,4.8));s=geo.groupby('Launch Site')['class'].agg(['mean','count']).sort_values('mean');plt.barh(s.index,100*s['mean']);plt.xlim(0,110);plt.xlabel('Landing success (%)');plt.title('Geolocated launch-site outcomes | 56-launch subset');[plt.text(r['mean']*100+1,i,'n='+str(int(r['count'])),va='center') for i,(_,r) in enumerate(s.iterrows())];plots[36]=save(36)
plt.figure(figsize=(11,4.8));g=geo.groupby('Launch Site').agg({'Lat':'first','Long':'first','class':'size'});plt.scatter(g['Long'],g['Lat'],s=g['class']*35,alpha=.6);[plt.annotate(idx,(r.Long,r.Lat),xytext=(5,7),textcoords='offset points',fontsize=10) for idx,r in g.iterrows()];plt.xlabel('Longitude');plt.ylabel('Latitude');plt.title('Launch-site locations (bubble area ∝ records)');plt.grid(alpha=.2);plots[37]=save(37)
plt.figure(figsize=(11,4.8));v=dash.groupby('Launch Site')['class'].sum().sort_values(ascending=False);plt.pie(v,labels=v.index,autopct='%1.0f%%',startangle=90);plt.title('Distribution of successful landings by launch site');plots[39]=save(39)
plt.figure(figsize=(11,4.8));
for k,g in dash.groupby('Launch Site'):plt.scatter(g['Payload Mass (kg)'],g['class'],label=k,alpha=.72,s=65)
plt.title('Landing outcome versus payload mass');plt.xlabel('Payload mass (kg)');plt.ylabel('Landing class (0 = no; 1 = yes)');plt.yticks([0,1]);plt.grid(alpha=.2);plt.legend(fontsize=10);plots[40]=save(40)
plt.figure(figsize=(11,4.8));q=dash.groupby('Launch Site')['class'].agg(['sum','count']);p=np.arange(len(q));plt.bar(p-.2,q['sum'],width=.4,label='Successful');plt.bar(p+.2,q['count']-q['sum'],width=.4,label='Other');plt.xticks(p,q.index,rotation=15);plt.ylabel('Flights');plt.title('Dashboard-style launch site comparison');plt.legend();plots[41]=save(41)
# independent reproducible model benchmarking (no test-set hyperparameter optimization)
X=a[['FlightNumber','PayloadMass','Flights','GridFins','Reused','Legs','Block','ReusedCount']].fillna(0).astype(float)
y=a.Class.astype(int);Xtr,Xte,ytr,yte=train_test_split(X,y,test_size=.2,random_state=42,stratify=y)
models={'Logistic regression':make_pipeline(StandardScaler(),LogisticRegression(max_iter=2000)),'SVM (RBF)':make_pipeline(StandardScaler(),SVC()),'Decision tree':DecisionTreeClassifier(max_depth=4,random_state=42),'KNN (k=5)':make_pipeline(StandardScaler(),KNeighborsClassifier(n_neighbors=5))}
scores={}; test_scores={}; trained={}
for name,m in models.items():
 cv=cross_val_score(m,Xtr,ytr,cv=StratifiedKFold(5,shuffle=True,random_state=42),scoring='accuracy');m.fit(Xtr,ytr);scores[name]=(cv.mean(),cv.std());test_scores[name]=accuracy_score(yte,m.predict(Xte));trained[name]=m
best=max(scores,key=lambda name:scores[name][0]); cm=confusion_matrix(yte,trained[best].predict(Xte),labels=[0,1]);
plt.figure(figsize=(11,4.8));names=list(scores);means=[scores[n][0]*100 for n in names];err=[scores[n][1]*100 for n in names];plt.bar(names,means,yerr=err,capsize=5);plt.ylim(0,110);plt.ylabel('5-fold cross-validated accuracy (%)');plt.xticks(rotation=12);plt.title('Model selection: training-fold CV accuracy (mean ± SD)');plots[43]=save(43)
plt.figure(figsize=(9,4.8));plt.imshow(cm,cmap='Blues');plt.title(f'Holdout confusion matrix | {best}');plt.xticks([0,1],['Predicted 0','Predicted 1']);plt.yticks([0,1],['Actual 0','Actual 1']);[plt.text(j,i,str(cm[i,j]),ha='center',va='center',fontsize=25,color='white' if cm[i,j]>cm.max()/2 else 'black') for i in (0,1) for j in (0,1)];plt.colorbar();plots[44]=save(44)
# SQL results provenance: use original database and genuine SQL queries
sql={24:"SELECT DISTINCT Launch_Site FROM SPACEXTBL",25:"SELECT Launch_Site, Date FROM SPACEXTBL WHERE Launch_Site LIKE 'CCA%' LIMIT 5",26:"SELECT SUM(PAYLOAD_MASS__KG_) AS Total_kg FROM SPACEXTBL WHERE Customer LIKE '%NASA%'",27:"SELECT AVG(PAYLOAD_MASS__KG_) AS Mean_kg FROM SPACEXTBL WHERE Booster_Version LIKE '%F9 v1.1%'",28:"SELECT MIN(Date) AS First_Date FROM SPACEXTBL WHERE Landing_Outcome LIKE '%Success%ground%'",29:"SELECT DISTINCT Booster_Version FROM SPACEXTBL WHERE Landing_Outcome LIKE '%Success%drone%' AND PAYLOAD_MASS__KG_ BETWEEN 4000 AND 6000",30:"SELECT Mission_Outcome, COUNT(*) AS N FROM SPACEXTBL GROUP BY Mission_Outcome",31:"SELECT Booster_Version, PAYLOAD_MASS__KG_ FROM SPACEXTBL ORDER BY PAYLOAD_MASS__KG_ DESC LIMIT 5",32:"SELECT Date, Booster_Version, Landing_Outcome FROM SPACEXTBL WHERE Landing_Outcome LIKE '%Failure%drone%' AND SUBSTR(Date,1,4)='2015'",33:"SELECT Landing_Outcome, COUNT(*) AS N FROM SPACEXTBL WHERE Date BETWEEN '2010-06-04' AND '2017-03-20' GROUP BY Landing_Outcome ORDER BY N DESC"}
queries={i:pd.read_sql_query(q,con) for i,q in sql.items()}
# Template manipulation
ppt=Presentation('/mnt/data/ds-capstone-template-coursera.pptx')
BLUE=RGBColor(0,70,157);DARK=RGBColor(35,48,69)
def set_shape(sh,text,size=19):
 sh.text=text
 for para in sh.text_frame.paragraphs:
  para.space_after=Pt(8)
  for run in para.runs:run.font.size=Pt(size);run.font.color.rgb=DARK

def clean(i):
 s=ppt.slides[i-1]
 # preserve slide title and numeric page footer; remove explanatory placeholders
 for sh in s.shapes:
  if not sh.has_text_frame or not sh.text.strip():continue
  txt=sh.text.strip(); top=sh.top/914400
  if txt==str(i) or (top<1.35 and len(txt)<90 and i not in (1,2,3,4,6,7,10,11,12,13,14,15,16,45,46,47)):continue
  if i==1 and ('<Name>' in txt or '<Date>' in txt):set_shape(sh,'M. Nicholas de Boulet\n9 October 2026',19)
  elif i!=1:sh.text=''
 return s

def textbox(s,text,x=.86,y=1.62,w=11.55,h=4.9,size=19):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 for j,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(12)
  for r in p.runs:r.font.size=Pt(size);r.font.color.rgb=DARK
 return sh

def picture(i,path,note=''):
 s=clean(i); s.shapes.add_picture(path,Inches(1.0),Inches(1.53),width=Inches(11.3),height=Inches(4.9))
 if note:textbox(s,note,1.0,6.48,11.2,.65,10)

def prose(i,txt,size=19):s=clean(i);textbox(s,txt,size=size)
prose(2,'Executive summary  •  Business problem and data sources\nMethodology  •  API, scraping, wrangling, SQL, EDA, geoanalytics, Dash, models\nResults  •  Observed data patterns, spatial context and classification\nConclusion  •  Findings, uncertainty, limitations and recommendations\nAppendix  •  Reproducibility notes and provenance',20)
prose(3,f'Objective: examine drivers of Falcon 9 first-stage landing outcomes and demonstrate reproducible analysis.\n90-launch analytical cohort (2010–2020); 60 successful landings, 30 other outcomes.\nObserved site-specific landing rates range from {a.groupby("LaunchSite").Class.mean().min()*100:.1f}% to {a.groupby("LaunchSite").Class.mean().max()*100:.1f}%.\nModel comparison uses stratified 5-fold validation on a training split; highest CV accuracy: {best} ({scores[best][0]*100:.1f}%).\nCaution: third-party reference archive; findings are independently reconstructed, not original course lab submissions.',18)
prose(4,'Commercial launch economics depend materially on recovering and reusing first-stage hardware.\nResearch question: which observed mission characteristics align with landing success, and how reliably can outcomes be predicted?\nScope: course-style historical samples, not a real-time operational forecast.\nAnalytic unit: one recorded first-stage launch; success is coded as Class = 1.',20)
prose(6,'1  DATA INGESTION   SpaceX API exercise and historical web-table records\n2  WRANGLING   Normalize fields, define binary landing class, inspect missingness\n3  EXPLORATION   Visual comparisons by launch site, orbit, payload, and time\n4  SPATIAL / DASHBOARD   Location plotting and interactive-style aggregations\n5  PREDICTION   Stratified training/holdout; cross-validated models\n6  AUDIT   Dataset and code provenance, caveats and reproducibility',18)
prose(7,'Reference material: published course-style SpaceX Falcon 9 datasets and notebooks.\nIndependent reconstruction uses the included CSV and SQLite assets; no live API requests are claimed.\nAnalytic sources: dataset_part_2.csv (90 records), SPACEXTBL (SQL), spacex_launch_geo.csv and spacex_launch_dash.csv (56-record subsets).\nDifferent extracts have different coverage and occasionally inconsistent fields; results are never merged without qualification.',18)
prose(8,'REPRODUCIBLE COLLECTION WORKFLOW\nSpaceX REST endpoints → JSON mission/rocket/launchpad records → field selection → tabular export.\nSample fields: flight number; date; booster; orbit; payload; site; landing outcome.\nThe provided CSV was independently inspected and analysed; original API notebook in the reference archive is not represented as user-authored.\nSource: course-style third-party reference archive; original SpaceX API documentation.',18)
prose(9,'REFERENCE WEB-SCRAPING WORKFLOW\nPublished launch-history table → HTML parsing → table rows → type normalization → structured CSV.\nData collected include booster version, mission outcome, customer, launch date and payload mass.\nThe archived web-scraped sample contains 121 records.\nProvenance: third-party reference notebook and spacex_web_scraped.csv; no personal scraping run claimed.',18)
prose(10,'Data integrity checks: record counts, column typing, category values, and missingness.\nTarget variable: Class = 1 for recorded successful first-stage landing; 0 otherwise.\n90 records; positive class 60 (66.7%), negative class 30 (33.3%).\nMissing LandingPad values (26) are not imputed as evidence of landing failure.\nFor modeling, eight numerical/binary predictors are used; missing numeric values receive zero only where applicable.\nSeparate SQL and 56-record spatial/dashboard sources retain their own denominator.',17)
prose(11,'Scatter plots: flight number and payload against site/orbit, coloured by landing outcome.\nBar charts: success rates by orbit; category labels include group size.\nLine series: annual share of successfully landed boosters from 2010–2020.\nInterpretation: temporal and site associations do not establish causality; unequal counts limit comparisons.\nCharts on slides 18–23 are independently generated from dataset_part_2.csv.',18)
prose(12,'SQLite inspection confirms SPACEXTBL table and landing/mission outcome fields.\nQueries use DISTINCT, LIKE, BETWEEN, MIN, AVG, SUM, GROUP BY, and ORDER BY.\nResults cover sites, NASA payload mass, booster version, ground/drone-ship outcomes and rankings.\nSQLite database contains historical date formatting quirks; treat ambiguous date results cautiously.\nSlides 24–33 show query-derived outputs, not screenshots of previously completed notebooks.',18)
prose(13,'Spatial reference: 56 geolocated launch records with latitude, longitude and landing class.\nUse coordinate markers to distinguish launches and locations.\nUse circles or symbol sizes to show launch frequency by site.\nCompare observed success share across sites; sample sizes must be displayed.\nSlides 35–37 are reproduced spatial analytical views, not captures of a live Folium session.',18)
prose(14,'Dashboard objective: filter flights by launch site and payload band to compare landing outcomes.\nSite success pie chart; payload-versus-outcome scatter; success-versus-other grouped comparison.\nA site selector and payload-mass range control are appropriate interactive elements.\nSlides 39–41 reproduce chart outputs from the 56-record Dash CSV; these are static analytical counterparts, not screenshots of a running app.',18)
prose(15,'Design: 90-record cohort; 80/20 stratified train/test split (random_state=42).\nPredictors: flight order, payload, flights, grid fins, reused, legs, block and reuse count.\nModels: regularized logistic regression, RBF SVM, shallow decision tree and k-nearest neighbours.\nModel choice: training-only 5-fold stratified cross-validation; holdout reserved for final assessment.\nLimitation: small, historical nonrandom sample; reported accuracies are illustrative, not deployable.',17)
prose(16,'EXPLORATORY RESULTS   Flight sequence, payload, orbit, launch site and annual trends\nSQL RESULTS   Ten transparent SQLite queries with reported outputs\nSPATIAL ANALYSIS   Launch-site coordinates and site-level comparisons\nDASHBOARD   Three independently reconstructed views\nPREDICTIVE ANALYSIS   Cross-validated model comparison and untouched holdout confusion matrix',19)
for i in list(range(18,24))+[35,36,37,39,40,41,43,44]: picture(i,plots[i],f'Source: third-party IBM-course reference CSV; independently reproduced analysis. n={len(a) if i<24 or i>=43 else len(geo)}. Not an original lab screenshot.')
for i,df in queries.items():
 vals=df.head(8).to_string(index=False,header=True)
 prose(i,f'SQL (independent rerun):\n{sql[i]}\n\nRESULTS ({len(df)} returned row(s)):\n{vals}\n\nSource: archived my_data1.db / SPACEXTBL; check date and classification semantics.',size=15 if len(vals)>190 else 17)
prose(45,f'1  Landing success in the 90-launch sample is {y.mean()*100:.1f}% (60/90); baseline accuracy is therefore already {max(y.mean(),1-y.mean())*100:.1f}%.\n2  Launch site, orbit, payload and time are associated with outcomes, but small subgroup samples constrain inference.\n3  Best training-fold model: {best} ({scores[best][0]*100:.1f}% CV); holdout accuracy {test_scores[best]*100:.1f}% (18 cases).\n4  Model reliability requires larger temporally held-out datasets, calibration and external validation.\n5  Responsible reporting preserves source provenance and differentiates reproduced analysis from prior lab work.',17)
prose(46,'DATA SOURCES\nThird-party IBM course-style reference archive: datasets, Jupyter notebooks and SQLite database supplied for this reconstruction.\nCourse grading template: IBM Skills Network, “Winning Space Race with Data Science,” 47-slide template.\nREPRODUCIBILITY\nAnalytic cohort: dataset_part_2.csv; static dashboard source: spacex_launch_dash.csv; locations: spacex_launch_geo.csv; SQL: my_data1.db.\nRepository created for this project: https://github.com/marvindebo/IBM-Applied-Data-Science-Capstone (upload analytical outputs separately).\nDISCLOSURE\nIllustrative independent reconstruction; not a claim that the associated hands-on labs were completed.',16)
# final blank slide as source / limitations
s=ppt.slides[46]; textbox(s,'Evidence and reproducibility',.85,.55,11,1,31);textbox(s,'Reference data: third-party SpaceX capstone ZIP supplied for independent analysis.\nOriginal course labs were not performed by the presenter; no experimental execution is implied.\nThe 90-record modeling extract and 56-record geospatial extract represent distinct sample frames.\nHistorical SQL archive contains date inconsistencies; date-based SQL findings are descriptive only.\nFor official submission, publish the independently generated notebooks/scripts and outputs in the project repository and check IBM’s academic-integrity rules.',.85,1.65,11.6,4.9,19)
# Decorative minor footer on analytical slides
for i in list(range(3,16))+list(range(18,34))+[35,36,37,39,40,41,43,44,45,46]:
 s=ppt.slides[i-1];tb=s.shapes.add_textbox(Inches(.65),Inches(7.15),Inches(9.5),Inches(.18)); p=tb.text_frame.paragraphs[0];p.text='Independent reconstruction • External reference data • 9 October 2026';p.runs[0].font.size=Pt(8);p.runs[0].font.color.rgb=RGBColor(105,119,142)
output=f'{out}/Data_Science_Capstone_Project_Report.pptx';ppt.save(output)
print('saved',output,'best',best,scores[best],test_scores[best], 'cm',cm.tolist()); print('sql rows',[(i,len(d)) for i,d in queries.items()]);print('charts',len(plots))
